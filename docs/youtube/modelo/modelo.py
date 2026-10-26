"""Matriz de nichos y modelo financiero a 12 meses del informe de YouTube.

Ejecutar: python3 docs/youtube/modelo/modelo.py
Genera: docs/youtube/modelo/salida.md (tablas que cita el informe)
        docs/youtube/Modelo_YouTube.xlsx (matriz, escenarios, sensibilidad)
Todas las cifras son simulaciones bajo hipótesis declaradas, no predicciones.
"""
import os, copy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

AQUI = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- matriz
PESOS = dict(D=.15, RPM=.15, G=.15, C=.10, P=.10, S=.10, E=.10, X=.10, R=.05)
# D demanda, RPM ingresos/visualización, G crecimiento/descubrimiento, C oportunidad competitiva,
# P facilidad/coste producción, S patrocinio/afiliación/otros, E evergreen, X externalización/escalado,
# R seguridad copyright/monetización.  Confianza: A buena evidencia, M media, B baja.
NICHOS = [
 ("IA aplicada al trabajo para autónomos y pymes (tutoriales y casos)", 8,7,8,7,7,9,5,6,9,"M"),
 ("Software de gestión para autónomos (facturación, Verifactu, banca)", 6,7,5,8,7,8,6,6,10,"M"),
 ("Finanzas personales con foco España (fiscalidad, vivienda, ahorro)", 9,9,7,3,5,9,7,6,6,"M"),
 ("Documentales de empresas españolas e hispanas", 7,6,7,8,5,6,9,8,7,"M"),
 ("Documentales de negocios globales en español (estilo MagnatesMedia)", 7,5,7,5,5,5,9,8,7,"M"),
 ("Historias de productos, marcas y fracasos empresariales", 7,5,7,6,5,5,9,8,7,"B"),
 ("Vivienda e inmobiliario en España explicado", 8,8,6,5,5,8,6,6,7,"B"),
 ("Tutoriales de ofimática y productividad (Excel, Notion, Google)", 8,6,6,4,6,8,8,6,10,"M"),
 ("Historia general e imperios en español", 8,3,7,4,5,4,10,8,7,"M"),
 ("Geografía y comparativas de países", 7,4,7,5,5,5,8,7,8,"B"),
 ("Psicología aplicada y comportamiento humano", 8,5,7,4,6,6,8,8,6,"B"),
 ("Megaconstrucciones, ingeniería e infraestructuras", 7,5,7,6,4,5,9,7,6,"B"),
 ("Habilidades profesionales y empleo (carrera, oposiciones, CV)", 7,5,5,6,6,7,7,5,9,"B"),
 ("Geopolítica y actualidad internacional", 8,5,8,3,5,4,4,6,5,"M"),
 ("Ciencia y curiosidades visuales", 8,4,7,4,3,4,9,5,7,"M"),
 ("Misterios y enigmas documentados", 8,2,8,4,6,3,8,8,5,"M"),
 ("True crime y casos judiciales documentados", 9,3,8,4,5,4,8,7,4,"M"),
 ("Historia de marcas de coches y automoción", 7,6,6,5,4,6,8,7,5,"B"),
 ("Biografías de personajes relevantes", 7,4,6,5,5,3,8,8,5,"B"),
 ("Lujo, grandes fortunas y alto patrimonio", 6,5,6,5,5,4,6,7,5,"B"),
 ("Viajes, destinos y análisis de costes", 7,5,5,5,4,7,5,6,6,"B"),
 ("Fórmula 1 y competición", 7,5,6,4,4,4,4,6,2,"M"),
 ("Análisis táctico e historias de fútbol", 8,3,6,4,4,4,3,6,2,"M"),
 ("Relatos narrados con IA / 'cash cow' genérico", 7,2,5,2,9,1,5,9,1,"M"),
]

def puntuar():
    filas = []
    for n in NICHOS:
        nombre, *s, conf = n
        d = dict(zip(PESOS, s))
        total = sum(d[k]*w for k, w in PESOS.items())
        filas.append((nombre, d, round(total, 2), conf))
    filas.sort(key=lambda f: -f[2])
    return filas

# ---------------------------------------------------------------- modelo
MESES = ["Nov-26","Dic-26","Ene-27","Feb-27","Mar-27","Abr-27","May-27","Jun-27","Jul-27","Ago-27","Sep-27","Oct-27"]

BASE = dict(
    nombre="Base",
    videos=[4,5,6,6,7,8,8,8,8,8,8,8],          # vídeos largos publicados/mes
    shorts=[8,10,12,12,12,12,12,12,12,12,12,12], # Shorts derivados (solo descubrimiento)
    v_inicial=250, v_final=3500,                 # vistas de un vídeo nuevo en su mes de estreno (m1 -> m12, crecimiento geométrico)
    cola=0.22, decaimiento=0.97,                 # catálogo: % de las vistas de estreno que conserva cada mes posterior
    vistas_short=[400,600,900,1200,1500,1800,2100,2400,2600,2800,3000,3200],  # vistas medias por Short
    avd_min=5.0,                                 # duración media de visualización (min) de vídeos largos
    subs_por_mil=14, subs_por_mil_short=0.6,     # suscriptores ganados por cada 1.000 vistas
    umbral_horas=8000, umbral_subs=1000,         # YPP para solicitudes desde el 1-feb-2027
    retraso_revision=1,                          # meses entre cumplir umbrales y cobrar
    rpm=3.5,                                     # € por 1.000 vistas largas una vez monetizado
    afil_por_mil=3.0, afil_desde=2,              # € afiliación por 1.000 vistas largas; mes en que empieza
    patro_desde=6, patro_cpm=18, patro_por_mes=[1,1,1,1,1,1,1,2,2,2,2,2],
    prod_desde=5, prod_conv=0.0002, prod_precio=39*0.9,  # producto digital: conversión sobre vistas, neto de comisiones
    servicios=[0,0,0,0,600,0,800,0,800,800,0,800],      # margen bruto de proyectos de implantación (opcional)
    herramientas=[55,55,55,110,110,110,120,120,120,120,120,120],
    edicion_por_video=[0,0,0,80,80,80,80,80,80,80,80,80],  # freelance de edición+miniatura desde el mes 4
    guion_por_video=[0]*12,
    otros=[100,10,10,10,150,150,150,150,150,150,150,150],  # m1: micrófono; desde m5: cuota autónomo tarifa plana + gestoría
)
CONSERVADOR = dict(BASE, nombre="Conservador",
    videos=[4,4,4,4,5,5,5,6,6,6,6,6],
    v_inicial=120, v_final=1300, cola=0.18, avd_min=4.0, subs_por_mil=10,
    vistas_short=[200,300,400,500,600,700,800,900,1000,1100,1200,1300],
    rpm=2.0, afil_por_mil=1.2, afil_desde=3, patro_desde=9, patro_cpm=12,
    patro_por_mes=[1]*12, prod_desde=7, prod_conv=0.0001,
    servicios=[0]*12,
    herramientas=[45,45,45,45,90,90,90,90,90,90,90,90],
    edicion_por_video=[0]*12,                    # sin tracción no se contrata: edita el propietario
    otros=[100]+[10]*11,                         # sin ingresos recurrentes no hay alta de autónomo en el modelo
)
OPTIMISTA = dict(BASE, nombre="Optimista",
    videos=[6,6,8,8,8,10,10,10,10,10,10,10],
    v_inicial=500, v_final=9000, cola=0.26, avd_min=6.0, subs_por_mil=18,
    vistas_short=[800,1200,1800,2500,3000,3500,4000,4500,5000,5500,6000,6500],
    rpm=5.5, afil_por_mil=5.0, afil_desde=2, patro_desde=4, patro_cpm=25,
    patro_por_mes=[1,1,1,2,2,2,3,3,3,3,4,4], prod_desde=4, prod_conv=0.0004,
    servicios=[0,0,800,800,1200,1200,1200,1200,1200,1200,1200,1200],
    herramientas=[70,70,130,130,150,150,150,150,150,150,150,150],
    edicion_por_video=[0,0,100,100,100,100,100,100,100,100,100,100],
    guion_por_video=[0]*5+[60]*7,
    otros=[100,10,150,150,150,150,150,150,150,150,150,150],
)

def simular(p, ajustes=None):
    p = copy.deepcopy(p)
    if ajustes: p.update(ajustes)
    n = 12
    g = (p["v_final"]/p["v_inicial"])**(1/(n-1))
    estreno = [p["v_inicial"]*g**i * p.get("factor_vistas",1) for i in range(n)]
    filas, cohortes, caja, horas_hist, subs = [], [], 0.0, [], 0
    mon_desde = None
    for m in range(n):
        nuevos = p["videos"][m]*estreno[m]
        catalogo = sum(c["vistas"]*p["cola"]*p["decaimiento"]**(m-c["mes"]-1)*c["n"] for c in cohortes)
        if p.get("sin_recomendacion_desde") is not None and m >= p["sin_recomendacion_desde"]:
            nuevos *= 0.35; catalogo *= 0.6
        extra_viral = p.get("viral",{}).get(m,0)
        largos = nuevos + catalogo + extra_viral
        if extra_viral: cohortes.append(dict(mes=m, vistas=extra_viral, n=1))
        cohortes.append(dict(mes=m, vistas=estreno[m], n=p["videos"][m]))
        shorts = p["shorts"][m]*p["vistas_short"][m]*p.get("factor_vistas",1)
        horas = largos*p["avd_min"]/60
        horas_hist.append(horas)
        subs += largos/1000*p["subs_por_mil"] + shorts/1000*p["subs_por_mil_short"]
        h12 = sum(horas_hist[-12:])
        if mon_desde is None and h12 >= p["umbral_horas"] and subs >= p["umbral_subs"]:
            mon_desde = max(m + p["retraso_revision"], p.get("monetizacion_minima",0))
        monetizado = mon_desde is not None and m >= mon_desde
        ads = largos/1000*p["rpm"]*p.get("factor_rpm",1) if monetizado else 0
        afil = largos/1000*p["afil_por_mil"] if m+1 >= p["afil_desde"] else 0
        patro = p["patro_por_mes"][m]*estreno[m]/1000*p["patro_cpm"] if m+1 >= p["patro_desde"] else 0
        patro = max(patro, 0)
        prod = largos*p["prod_conv"]*p["prod_precio"] if m+1 >= p["prod_desde"] else 0
        serv = p["servicios"][m]
        ingresos = ads+afil+patro+prod+serv
        f = p.get("factor_coste",1)
        c_herr = p["herramientas"][m]
        c_ed = p["videos"][m]*(p["edicion_por_video"][m]+p["guion_por_video"][m])*f
        c_otros = p["otros"][m]
        costes = c_herr+c_ed+c_otros
        bai = ingresos-costes
        caja += bai
        filas.append(dict(mes=MESES[m], videos=p["videos"][m], shorts=p["shorts"][m],
            v_nuevos=nuevos+extra_viral, v_catalogo=catalogo, v_largos=largos, v_shorts=shorts,
            horas12=h12, subs=subs, estado="Monetizado" if monetizado else ("Elegible-revisión" if mon_desde is not None else "No elegible"),
            ads=ads, afil=afil, patro=patro, prod=prod, serv=serv, ingresos=ingresos,
            c_herr=c_herr, c_ed=c_ed, c_otros=c_otros, costes=costes, bai=bai, caja=caja))
    return filas

def resumen(filas):
    caja_min = min(f["caja"] for f in filas)
    eq = next((f["mes"] for f in filas if f["bai"] > 0 and f["ingresos"]-f["serv"] > f["costes"]), "no en 12 meses")
    return dict(ingresos=sum(f["ingresos"] for f in filas), costes=sum(f["costes"] for f in filas),
                bai=sum(f["bai"] for f in filas), caja_min=caja_min, caja_final=filas[-1]["caja"],
                equilibrio_sin_servicios=eq,
                mon=next((f["mes"] for f in filas if f["estado"]=="Monetizado"), "no en 12 meses"),
                ingreso_m12=filas[-1]["ingresos"], bai_m12=filas[-1]["bai"],
                vistas_m12=filas[-1]["v_largos"], subs_m12=filas[-1]["subs"])

def e(x): return f"{x:,.0f}".replace(",", ".")

def tabla_escenario(filas):
    cab = "| Mes | Vídeos | Shorts | Vistas estreno | Vistas catálogo | Vistas largos | Vistas Shorts | Horas 12m | Subs | Estado YPP | Anuncios € | Afiliación € | Patrocinio € | Producto € | Servicios € | Ingresos € | Herram. € | Edición € | Otros € | BAI € | Caja acum. € |\n"
    cab += "|"+"---|"*21+"\n"
    for f in filas:
        cab += f"| {f['mes']} | {f['videos']} | {f['shorts']} | {e(f['v_nuevos'])} | {e(f['v_catalogo'])} | {e(f['v_largos'])} | {e(f['v_shorts'])} | {e(f['horas12'])} | {e(f['subs'])} | {f['estado']} | {e(f['ads'])} | {e(f['afil'])} | {e(f['patro'])} | {e(f['prod'])} | {e(f['serv'])} | {e(f['ingresos'])} | {e(f['c_herr'])} | {e(f['c_ed'])} | {e(f['c_otros'])} | {e(f['bai'])} | {e(f['caja'])} |\n"
    return cab

def main():
    out = []
    filas = puntuar()
    out.append("## Matriz de nichos (ordenada)\n")
    out.append("| # | Nicho | Dem 15% | RPM 15% | Crec 15% | Comp 10% | Prod 10% | Patro/Afil 10% | Evergr 10% | Escala 10% | Seguridad 5% | **Total** | Confianza |")
    out.append("|"+"---|"*13)
    for i,(n,d,t,c) in enumerate(filas,1):
        out.append(f"| {i} | {n} | " + " | ".join(str(d[k]) for k in PESOS) + f" | **{t:.2f}** | {c} |")
    escenarios = [CONSERVADOR, BASE, OPTIMISTA]
    res = {}
    for p in escenarios:
        fs = simular(p); res[p["nombre"]] = (fs, resumen(fs))
        out.append(f"\n## Escenario {p['nombre']}\n")
        out.append(tabla_escenario(fs))
        r = resumen(fs)
        out.append(f"Resumen: ingresos 12m {e(r['ingresos'])} €, costes {e(r['costes'])} €, BAI {e(r['bai'])} €, caja mínima {e(r['caja_min'])} €, caja final {e(r['caja_final'])} €, monetización YPP desde {r['mon']}, primer mes con beneficio sin contar servicios: {r['equilibrio_sin_servicios']}, ingresos mes 12 {e(r['ingreso_m12'])} €, BAI mes 12 {e(r['bai_m12'])} €, vistas largas mes 12 {e(r['vistas_m12'])}, suscriptores mes 12 {e(r['subs_m12'])}.\n")
    # sensibilidad sobre el escenario base
    sens = [
        ("Base", {}),
        ("Visualizaciones −50 %", dict(factor_vistas=0.5)),
        ("RPM a la mitad", dict(factor_rpm=0.5)),
        ("Producción al doble de coste", dict(factor_coste=2.0)),
        ("Monetización YPP 6 meses más tarde (fuera del año 1)", dict(monetizacion_minima=16)),
        ("Referencia: umbral antiguo de 4.000 h (no aplicable a solicitudes desde 1-feb-2027)", dict(umbral_horas=4000)),
        ("Un vídeo viral (+250.000 vistas en el mes 6)", dict(viral={5:250000})),
        ("Pérdida de recomendaciones desde el mes 8", dict(sin_recomendacion_desde=7)),
        ("Base sin servicios", dict(servicios=[0]*12)),
    ]
    out.append("\n## Sensibilidad (sobre el escenario base)\n")
    out.append("| Caso | Ingresos 12m € | Costes 12m € | BAI 12m € | Caja mínima € | Mes de monetización YPP | Ingresos mes 12 € | BAI mes 12 € |")
    out.append("|---|---|---|---|---|---|---|---|")
    sens_res = []
    for nombre, aj in sens:
        r = resumen(simular(BASE, aj)); sens_res.append((nombre, r))
        out.append(f"| {nombre} | {e(r['ingresos'])} | {e(r['costes'])} | {e(r['bai'])} | {e(r['caja_min'])} | {r['mon']} | {e(r['ingreso_m12'])} | {e(r['bai_m12'])} |")
    out.append("\n## Análisis inverso\n")
    mezclas = [
        ("Solo AdSense, RPM 3,5 €", 3.5, 0),
        ("Solo AdSense, RPM 2 € (audiencia más latinoamericana)", 2.0, 0),
        ("Mezcla madura: anuncios 3,5 + afiliación 3 + patrocinio 2 + producto 7 = 15,5 € por mil", 15.5, 0),
        ("Mezcla madura + 1 proyecto de servicios de 1.500 € de margen al mes", 15.5, 1500),
    ]
    out.append("| Mezcla de ingresos | 1.000 €/mes | 3.000 €/mes | 5.000 €/mes | 10.000 €/mes |")
    out.append("|---|---|---|---|---|")
    for nombre, por_mil, fijo in mezclas:
        celdas = []
        for obj in (1000,3000,5000,10000):
            v = max(obj-fijo,0)/por_mil*1000
            celdas.append(e(v)+" vistas")
        out.append(f"| {nombre} | " + " | ".join(celdas) + " |")
    out.append("\nBeneficio operativo objetivo (estructura de costes por tramo):\n")
    out.append("| Beneficio objetivo | Vídeos/mes | Coste por vídeo | Herramientas | Costes totales | Ingresos necesarios | Vistas largas/mes necesarias a 15,5 €/mil |")
    out.append("|---|---|---|---|---|---|---|")
    for obj, vids, cpv, herr in ((1000,8,80,120),(3000,10,150,200),(5000,12,220,300),(10000,16,300,450)):
        costes = vids*cpv+herr+150
        ing = obj+costes
        out.append(f"| {e(obj)} € | {vids} | {cpv} € | {herr} € | {e(costes)} € | {e(ing)} € | {e(ing/15.5*1000)} |")
    with open(os.path.join(AQUI, "salida.md"), "w") as fh:
        fh.write("\n".join(out)+"\n")

    # ---------------------------------------------------------- xlsx
    wb = Workbook(); ws = wb.active; ws.title = "Matriz nichos"
    bold = Font(bold=True); head = PatternFill("solid", fgColor="1F3A5F"); white = Font(bold=True, color="FFFFFF")
    hdr = ["Nicho"]+[f"{k} ({int(w*100)}%)" for k,w in PESOS.items()]+["Total (fórmula)","Confianza"]
    ws.append(hdr)
    for i,(n,d,t,c) in enumerate(filas, start=2):
        ws.append([n]+[d[k] for k in PESOS]+[None,c])
        cols = "BCDEFGHIJ"
        ws[f"K{i}"] = "=" + "+".join(f"{cols[j]}{i}*{w}" for j,w in enumerate(PESOS.values()))
    for cell in ws[1]: cell.font = white; cell.fill = head; cell.alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 62
    for p in escenarios:
        fs = res[p["nombre"]][0]
        w = wb.create_sheet(p["nombre"])
        keys = ["mes","videos","shorts","v_nuevos","v_catalogo","v_largos","v_shorts","horas12","subs","estado","ads","afil","patro","prod","serv","ingresos","c_herr","c_ed","c_otros","costes","bai","caja"]
        w.append(["Mes","Vídeos largos","Shorts","Vistas vídeos nuevos","Vistas catálogo","Vistas largos total","Vistas Shorts","Horas públicas 12m","Suscriptores","Estado YPP","Anuncios €","Afiliación €","Patrocinio €","Producto digital €","Servicios €","Ingresos €","Herramientas €","Edición/colaboradores €","Otros €","Costes €","BAI €","Caja acumulada €"])
        for f in fs:
            w.append([round(f[k],0) if isinstance(f[k], float) else f[k] for k in keys])
        for r in range(2, 14):
            w[f"P{r}"] = f"=K{r}+L{r}+M{r}+N{r}+O{r}"
            w[f"T{r}"] = f"=Q{r}+R{r}+S{r}"
            w[f"U{r}"] = f"=P{r}-T{r}"
            w[f"V{r}"] = f"=U{r}" if r == 2 else f"=V{r-1}+U{r}"
        w.append([]); w.append(["Hipótesis"])
        for k,v in p.items():
            w.append([k, str(v)])
        for cell in w[1]: cell.font = white; cell.fill = head; cell.alignment = Alignment(wrap_text=True)
    w = wb.create_sheet("Sensibilidad")
    w.append(["Caso","Ingresos 12m €","Costes 12m €","BAI 12m €","Caja mínima €","Mes monetización","Ingresos mes 12 €","BAI mes 12 €"])
    for nombre, r in sens_res:
        w.append([nombre, round(r["ingresos"]), round(r["costes"]), round(r["bai"]), round(r["caja_min"]), r["mon"], round(r["ingreso_m12"]), round(r["bai_m12"])])
    for cell in w[1]: cell.font = white; cell.fill = head
    w.column_dimensions["A"].width = 60
    import ideas
    w = wb.create_sheet("Banco 100 ideas")
    w.append(["#","Categoría","Título","Concepto","Por qué interesa","Gancho","Estructura","Recursos","Dificultad","Diferenciación","Riesgos","Descubrimiento","Evergreen","Facilidad","Coste €","Intención","Competencia","Comercial","Encaje"])
    k = 0
    for cat, lista in ideas.CATEGORIAS:
        for i in lista:
            k += 1; w.append([k, cat, *i])
    for cell in w[1]: cell.font = white; cell.fill = head
    w.column_dimensions["C"].width = 60
    wb.save(os.path.join(os.path.dirname(AQUI), "Modelo_YouTube.xlsx"))

if __name__ == "__main__":
    main()
